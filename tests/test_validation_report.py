"""
src/validation_report.py and the (6, 1) suite of src/validate_constraints.py: edge cases, scoring and the Excel
workbooks, on synthetic rebuild results (no SolidWorks).
"""

import json

import numpy as np
import pandas as pd
import pytest

import constraints as C
import validate_constraints as V
import validation_report as R

OD, L, G = 40.0, 50.0, 3.0
CFG = (6, 1)


def accepted(df):
    return C.feasible_mask(df.assign(N=6, T=1), OD, L, G).to_numpy()


def write_run(tmp_path, pts, built, n_meta=None):
    """swgen_results.jsonl + meta CSV for `pts` (DataFrame with R, A, CC, VC, source, label)."""
    results = tmp_path / "swgen_results.jsonl"
    with open(results, "w") as f:
        for i, (rec, ok) in enumerate(zip(pts.to_dict("records"), built), start=1):
            f.write(json.dumps({"row": i, "params": {k: rec[k] for k in V.PARAM_COLS}, "status": "ok" if ok else "failed",
                                "rebuild_errors": [] if ok else [{"feature": "Wrap2", "code": 1}]}) + "\n")
    meta = pts.copy()
    meta.insert(0, "row", range(1, len(meta) + 1))
    meta_path = tmp_path / "points.meta.csv"
    meta.to_csv(meta_path, index=False)
    return str(results), str(meta_path)


class TestEdgeCases:
    def test_suite_has_unique_in_box_points_and_is_reproducible(self):
        a = V.suite_points(CFG, OD, L, G, n_sweep=120, seed=7)
        b = V.suite_points(CFG, OD, L, G, n_sweep=120, seed=7)
        pd.testing.assert_frame_equal(a, b)
        assert (a["source"] == "sweep").sum() == 120 and 70 <= (a["source"] == "edge").sum() <= 120
        assert not a.duplicated(V.PARAM_COLS).any()
        for col, (lo, hi) in C.PARAM_BOUNDS.items():
            assert a[col].between(lo, hi).all(), col
        assert (a["R"].round(2) == a["R"]).all() and (a["A"].round(2) == a["A"]).all()

    def test_only_6_1_has_edge_cases(self):
        with pytest.raises(NotImplementedError):
            V.suite_points((6, 0), OD, L, G)

    def test_brackets_straddle_their_boundary(self):
        edge = V.edge_points_61(OD, L, G)
        ok = accepted(edge)
        kind = edge["label"].str.extract(r"^(box corner|C1|C5|C4' tilt|Vsep=2R) ?([-+]1%)?")
        assert (kind[0] == "box corner").sum() == 16
        for name in ("C1", "C5", "C4' tilt"):
            low = (kind[0] == name) & (kind[1] == "-1%")
            high = (kind[0] == name) & (kind[1] == "+1%")
            # a bracket whose +1 % side would leave the design box (R > 8.8) keeps only its -1 % side
            assert low.sum() >= 5 and 0 <= low.sum() - high.sum() <= 1, name
            low, high = low & edge["label"].isin(edge.loc[high, "label"].str.replace("+1%", "-1%", regex=False)), high
            inside, outside = (ok[low], ok[high]) if name != "C4' tilt" else (ok[high], ok[low])   # tilt: more A is inside
            assert inside.all() and not outside.any(), name

    def test_each_bracket_is_decided_by_its_own_constraint(self):
        """For a C1 / C5 / C4' bracket the other two constraints hold on both sides, so only the named one flips."""
        edge = V.edge_points_61(OD, L, G)
        rep = C.evaluate(edge.assign(N=6, T=1), OD, L, G)
        for prefix, own in (("C1 ", "C1"), ("C5 ", "C5"), ("C4' tilt", "C4'")):
            sel = edge["label"].str.startswith(prefix).to_numpy()
            others = [n for n in ("C1", "C4'", "C5") if n != own]
            assert rep.loc[sel, others].astype(bool).all().all(), prefix
            assert rep.loc[sel, own].astype(bool).nunique() == 2, prefix


class TestScoring:
    def test_outcomes_and_confusion(self):
        assert R.outcome_of([1, 1, 0, 0], [1, 0, 1, 0]).tolist() == ["TP", "FP", "FN", "TN"]
        c = R.confusion(["TP", "TP", "TP", "FP", "FN", "TN", "TN"])
        assert (c["TP"], c["FP"], c["FN"], c["TN"], c["n"]) == (3, 1, 1, 2, 7)
        assert c["accuracy"] == pytest.approx(5 / 7, abs=1e-4) and c["precision"] == 0.75
        assert c["recall"] == 0.75 and c["specificity"] == pytest.approx(2 / 3, abs=1e-4)
        assert R.confusion([])["accuracy"] is None and R.confusion(["TP"])["specificity"] is None

    def test_constraints_that_match_reality_have_no_errors(self):
        pts = V.random_box_points(500, seed=3)
        scored = R.score_points(pts.assign(built=accepted(pts)), CFG, OD, L, G)
        assert set(scored["outcome"]) <= {"TP", "TN"} and scored["accepted"].sum() == (scored["outcome"] == "TP").sum() > 0
        assert {"C1", "C4'", "C5"} <= set(scored.columns)

    def test_per_constraint_table_counts_add_up(self):
        pts = V.random_box_points(400, seed=4)
        built = accepted(pts) & (pts["A"].to_numpy() > 40)               # reality needs A > 40 as well
        scored = R.score_points(pts.assign(built=built), CFG, OD, L, G)
        t = R.per_constraint_table(scored, CFG)
        assert list(t["constraint"]) == ["C1", "C4'", "C5"]
        assert (t[["TP", "FP", "FN", "TN"]].sum(axis=1) == 400).all()
        assert (t["FN"] == t["violated_by_built"]).all() and (t["FP"] == t["satisfied_by_failed"]).all()
        assert R.summary_table(scored).loc[0, "n"] == 400


class TestWorkbooks:
    def test_run_workbook_lists_every_point_and_the_constraints(self, tmp_path):
        from openpyxl import load_workbook
        suite = V.suite_points(CFG, OD, L, G, n_sweep=40, seed=5)
        built = accepted(suite) | (np.arange(len(suite)) % 7 == 0)        # some false negatives... and positives
        built &= ~(np.arange(len(suite)) % 11 == 0)                       # ... and some false positives
        results, meta = write_run(tmp_path, suite, built)
        scored = R.score_points(R.load_run(results, meta), CFG, OD, L, G)
        out = tmp_path / "report.xlsx"
        tables = R.write_run_workbook(str(out), scored, CFG, OD, L, G, "title")

        wb = load_workbook(out)
        assert wb.sheetnames == ["Constraints", "Summary", "All points", "True positives", "False positives",
                                 "False negatives", "True negatives"]
        text = " ".join(str(c.value) for row in wb["Constraints"].iter_rows() for c in row if c.value)
        assert "C1" in text and "C4'" in text and "C5" in text and "Vsep" in text
        assert wb["All points"].max_row == len(suite) + 1
        counts = scored["outcome"].value_counts()
        for code, sheet in (("TP", "True positives"), ("FP", "False positives"), ("FN", "False negatives"), ("TN", "True negatives")):
            assert counts.get(code, 0) > 0, code                        # premise of this test
            assert wb[sheet].max_row == counts[code] + 1, sheet
        all_row = tables["summary"].iloc[0]
        assert all_row["n"] == len(suite) and all_row["TP"] + all_row["FP"] + all_row["FN"] + all_row["TN"] == len(suite)
        assert list(tables["summary"]["subset"]) == ["All points", "edge only", "sweep only"]

    def test_unfinished_run_is_refused_unless_allowed(self, tmp_path):
        suite = V.suite_points(CFG, OD, L, G, n_sweep=10, seed=6)
        results, meta = write_run(tmp_path, suite.iloc[:5], [True] * 5)
        pd.concat([pd.read_csv(meta).iloc[:5]]).to_csv(meta, index=False)
        full_meta = suite.copy()
        full_meta.insert(0, "row", range(1, len(suite) + 1))
        full_meta.to_csv(meta, index=False)
        with pytest.raises(ValueError, match="not finished"):
            R.load_run(results, meta)
        assert len(R.load_run(results, meta, allow_partial=True)) == 5

    def test_results_must_match_the_requested_points(self, tmp_path):
        suite = V.suite_points(CFG, OD, L, G, n_sweep=10, seed=6)
        results, meta = write_run(tmp_path, suite, [True] * len(suite))
        m = pd.read_csv(meta)
        m.loc[3, "R"] += 0.5
        m.to_csv(meta, index=False)
        with pytest.raises(ValueError, match="do not match"):
            R.load_run(results, meta)

    def test_old_run_comparison(self, tmp_path, capsys):
        from openpyxl import load_workbook
        pts = V.random_box_points(120, seed=8)
        pts["source"] = np.where(np.arange(120) % 2 == 0, "sweep", "edge")
        pts["label"] = "x"
        pts["built"] = accepted(pts) | (pts["A"].to_numpy() > 85)
        pts["current_spec_accepts"] = accepted(pts) & (pts["A"].to_numpy() > 60)        # the "old" verdicts
        csv = tmp_path / "old.csv"
        pts.to_csv(csv, index=False)
        out = tmp_path / "old.xlsx"
        assert R.main(["old", "--old-csv", str(csv), "--config", "6:1", "--OD", "40", "--L", "50", "--G", "3", "--out", str(out)]) == 0
        wb = load_workbook(out)
        assert wb.sheetnames == ["Constraints", "Summary", "All points", "Changed verdicts", "False positives (new)", "False negatives (new)"]
        assert wb["All points"].max_row == 121
        header = [c.value for c in wb["All points"][1]]
        assert {"old_accepts", "old_outcome", "new_outcome", "change"} <= set(header)
        changed = int(((pts["current_spec_accepts"]) != accepted(pts)).sum())
        assert changed > 0 and wb["Changed verdicts"].max_row == changed + 1
        assert "workbook written" in capsys.readouterr().out

    def test_cli_errors(self, tmp_path):
        bad = tmp_path / "bad.csv"
        pd.DataFrame({"R": [3.0]}).to_csv(bad, index=False)
        assert R.main(["old", "--old-csv", str(bad), "--config", "6:1", "--OD", "40", "--L", "50", "--out", str(tmp_path / "x.xlsx")]) == 2
        assert R.main(["old", "--old-csv", str(bad), "--config", "0:1", "--OD", "40", "--L", "50", "--out", str(tmp_path / "x.xlsx")]) == 2

    def test_suite_cli_writes_points_and_meta(self, tmp_path):
        out = tmp_path / "suite.csv"
        assert V.main(["suite", "--config", "6:1", "--OD", "40", "--L", "50", "--G", "3", "--n-sweep", "30", "--seed", "1", "--out", str(out)]) == 0
        pts, meta = pd.read_csv(out), pd.read_csv(tmp_path / "suite.meta.csv")
        assert list(pts.columns) == ["R", "A", "CC", "VC"] and len(pts) == len(meta)
        assert meta["row"].tolist() == list(range(1, len(meta) + 1)) and set(meta["source"]) == {"edge", "sweep"}
        assert V.main(["suite", "--config", "6:0", "--OD", "40", "--L", "50", "--out", str(out)]) == 2
