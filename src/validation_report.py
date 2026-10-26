"""
Excel reports comparing the constraints of a configuration with SolidWorks rebuild results.

Ground truth is the rebuild (SwGen status ok / warning = built, anything else = failed).  A point is a

    TP  true positive    constraints accept it,  SolidWorks built it
    FP  false positive   constraints accept it,  SolidWorks could NOT build it   (constraints too loose)
    FN  false negative   constraints reject it,  SolidWorks built it              (constraints too strict)
    TN  true negative    constraints reject it,  SolidWorks could not build it

Two workbooks:

  run   a fresh SwGen run (points from `validate_constraints.py suite`, results = swgen_results.jsonl)
  old   an earlier run that is already stored as a CSV (R, A, CC, VC, built, ...) rescored against the
        CURRENT constraints, next to the verdict the constraints of that time gave (current_spec_accepts)

Usage:

    python src/validation_report.py run --results out/swgen_results.jsonl --meta points.meta.csv \\
           --config 6:1 --OD 40 --L 50 --G 3 --out report.xlsx
    python src/validation_report.py old --old-csv validation_results/points_N6TAShell_6_1.csv \\
           --config 6:1 --OD 40 --L 50 --G 3 --out report_old.xlsx
"""

from __future__ import annotations

import argparse
import sys

import numpy as np
import pandas as pd

from constraints import PARAM_COLS, constraints_for, evaluate, parse_configs
from validate_constraints import load_results

OUTCOMES = ("TP", "FP", "FN", "TN")

DESCRIPTIONS = {
    "C1": "shared with (6, 0)",
    "C2": "(6, 0) only",
    "C3": "(6, 0) only",
    "C4'": "twist-angle condition with Vsep = (L - 2(R+G))/(VC - 1); replaces the old minimum-tilt C4",
    "C5": "vertical packing; replaces C3 (equals C3 at A = 90 deg)",
}

OLD_SPEC_TEXT = (
    "Constraints in force when the old points were rebuilt: C1, C2, C3 as (6, 0) plus the minimum tilt "
    "C4: A > arctan((v + 2R)/((4*sqrt(3)/3)*v - 2*sqrt(3)*R)), v = 3*pi*OD/(2*sqrt(3)*CC)."
)


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------

def outcome_of(accepted, built) -> np.ndarray:
    accepted, built = np.asarray(accepted, dtype=bool), np.asarray(built, dtype=bool)
    return np.select([accepted & built, accepted & ~built, ~accepted & built], ["TP", "FP", "FN"], default="TN")


def score_points(df: pd.DataFrame, config: tuple[int, int], OD: float, L: float, G: float = 0.0) -> pd.DataFrame:
    """`df` needs R, A, CC, VC and a boolean `built`; every other column is carried along.  Adds one boolean
    column per constraint of the configuration, `accepted` (all constraints hold) and `outcome` (TP/FP/FN/TN)."""
    df = df.reset_index(drop=True)
    report = evaluate(df.assign(N=config[0], T=config[1]), OD, L, G)
    out = df.copy()
    for con in constraints_for(config):
        out[con.name] = report[con.name].astype(bool).to_numpy()
    out["accepted"] = report["feasible"].to_numpy(dtype=bool)
    out["built"] = df["built"].astype(bool).to_numpy()
    out["outcome"] = outcome_of(out["accepted"], out["built"])
    return out


def confusion(outcomes) -> dict:
    """Counts and rates from a sequence of TP/FP/FN/TN labels."""
    o = pd.Series(list(outcomes), dtype="object")
    c = {k: int((o == k).sum()) for k in OUTCOMES}
    n = sum(c.values())

    def ratio(a, b):
        return round(a / b, 4) if b else None

    return {"n": n, "built": c["TP"] + c["FN"], "failed": c["FP"] + c["TN"], **c,
            "accuracy": ratio(c["TP"] + c["TN"], n),
            "precision": ratio(c["TP"], c["TP"] + c["FP"]),
            "recall": ratio(c["TP"], c["TP"] + c["FN"]),
            "specificity": ratio(c["TN"], c["TN"] + c["FP"])}


def subsets(scored: pd.DataFrame) -> list[tuple[str, pd.Series]]:
    """('All points', mask) followed by one entry per `source` value, if the frame has a source column."""
    out = [("All points", pd.Series(True, index=scored.index))]
    if "source" in scored.columns:
        for src in scored["source"].dropna().unique():
            out.append((f"{src} only", scored["source"] == src))
    return out


def summary_table(scored: pd.DataFrame, outcome_col: str = "outcome", spec: str | None = None) -> pd.DataFrame:
    rows = []
    for name, mask in subsets(scored):
        row = {"subset": name, **confusion(scored.loc[mask, outcome_col])}
        if spec is not None:
            row = {"spec": spec, **row}
        rows.append(row)
    return pd.DataFrame(rows)


def per_constraint_table(scored: pd.DataFrame, config: tuple[int, int]) -> pd.DataFrame:
    """Each constraint on its own as a predictor of "built", plus the three diagnostics of validate_constraints."""
    cset = constraints_for(config)
    built = scored["built"].to_numpy(dtype=bool)
    rows = []
    for con in cset:
        sat = scored[con.name].to_numpy(dtype=bool)
        others = np.ones(len(scored), dtype=bool)
        for other in cset:
            if other.name != con.name:
                others &= scored[other.name].to_numpy(dtype=bool)
        rows.append({
            "constraint": con.name, "expression": con.expr,
            **{k: v for k, v in confusion(outcome_of(sat, built)).items() if k in ("TP", "FP", "FN", "TN")},
            "violated_by_built": int((~sat & built).sum()),
            "satisfied_by_failed": int((sat & ~built).sum()),
            "only_violation_of_failed": int((~sat & others & ~built).sum()),
        })
    return pd.DataFrame(rows)


def constraints_table(config: tuple[int, int]) -> pd.DataFrame:
    return pd.DataFrame([{"constraint": c.name, "expression (strict inequalities)": c.expr,
                          "note": DESCRIPTIONS.get(c.name, "")} for c in constraints_for(config)])


# ---------------------------------------------------------------------------
# Workbook writing
# ---------------------------------------------------------------------------

_FILLS = {"FP": "F8CBAD", "FN": "FFE699", "TP": "C6E0B4", "TN": "DDEBF7"}


def _style_sheet(ws, header_row: int = 1, outcome_col: str | None = None, freeze: bool = True) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="D9D9D9")
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 70)
    if freeze:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if outcome_col:
        idx = next((c.column for c in ws[header_row] if c.value == outcome_col), None)
        if idx:
            for row in ws.iter_rows(min_row=header_row + 1):
                code = row[idx - 1].value
                if code in _FILLS:
                    row[idx - 1].fill = PatternFill("solid", fgColor=_FILLS[code])


def _write_block(ws, df: pd.DataFrame, row: int, title: str | None = None) -> int:
    """Write an optional bold title and a DataFrame at `row`; return the next free row."""
    from openpyxl.styles import Font, PatternFill

    if title:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
        row += 1
    for j, name in enumerate(df.columns, start=1):
        c = ws.cell(row=row, column=j, value=str(name))
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9D9D9")
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(rec, start=1):
            if isinstance(v, (np.bool_, bool)):
                v = bool(v)
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.floating):
                v = None if np.isnan(v) else float(v)
            elif v is pd.NA:
                v = None
            ws.cell(row=row + i, column=j, value=v)
    return row + len(df) + 2


def _data_sheet(wb, name: str, df: pd.DataFrame, empty_note: str) -> None:
    ws = wb.create_sheet(name)
    if df.empty:
        ws.cell(row=1, column=1, value=empty_note)
        return
    _write_block(ws, df, 1)
    _style_sheet(ws, 1, outcome_col="outcome" if "outcome" in df.columns else None)


def _legend(extra: list[str] | None = None) -> pd.DataFrame:
    rows = [
        ("Positive", "the constraints ACCEPT the point (all of them hold)"),
        ("Built", "SolidWorks rebuild status ok or warning (ground truth)"),
        ("TP true positive", "accepted and built"),
        ("FP false positive", "accepted but SolidWorks could NOT build it (constraints too loose)"),
        ("FN false negative", "rejected although SolidWorks built it (constraints too strict)"),
        ("TN true negative", "rejected and SolidWorks could not build it"),
        ("accuracy", "(TP + TN) / n"), ("precision", "TP / (TP + FP)"),
        ("recall", "TP / (TP + FN)"), ("specificity", "TN / (TN + FP)"),
    ] + [("", t) for t in (extra or [])]
    return pd.DataFrame(rows, columns=["term", "meaning"])


def _point_columns(scored: pd.DataFrame, names: list[str], extra_front: list[str], extra_back: list[str]) -> pd.DataFrame:
    cols = [c for c in extra_front if c in scored.columns] + PARAM_COLS + names + \
           [c for c in extra_back if c in scored.columns]
    return scored[cols]


def write_run_workbook(path: str, scored: pd.DataFrame, config: tuple[int, int], OD: float, L: float, G: float,
                       title: str) -> dict:
    """Constraints / Summary / All points / TP / FP / FN / TN sheets for a fresh run."""
    from openpyxl import Workbook

    names = [c.name for c in constraints_for(config)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Constraints"
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=2, column=1, value=f"Configuration (N={config[0]}, T={config[1]});  OD={OD:g}  L={L:g}  G={G:g};  "
                                   f"A in degrees;  Vsep = (L - 2(R+G))/(VC - 1)")
    _write_block(ws, constraints_table(config), 4, "Constraints tested")
    _style_sheet(ws, 5, freeze=False)

    ws = wb.create_sheet("Summary")
    summary = summary_table(scored)
    row = _write_block(ws, summary, 1, "Confusion matrix (constraints vs SolidWorks rebuild)")
    row = _write_block(ws, per_constraint_table(scored, config), row, "Each constraint on its own")
    _write_block(ws, _legend(), row, "Legend")
    _style_sheet(ws, 2, freeze=False)

    front = ["row", "source", "label"]
    back = ["accepted", "built", "status", "failed_features", "outcome"]
    table = _point_columns(scored, names, front, back).rename(columns={"accepted": "constraints_accept", "built": "solidworks_built"})
    _data_sheet(wb, "All points", table, "no points")
    for code, sheet in (("TP", "True positives"), ("FP", "False positives"), ("FN", "False negatives"), ("TN", "True negatives")):
        _data_sheet(wb, sheet, table[scored["outcome"] == code], f"No {sheet.lower()}.")
    wb.save(path)
    return {"summary": summary, "per_constraint": per_constraint_table(scored, config)}


def write_comparison_workbook(path: str, scored: pd.DataFrame, config: tuple[int, int], OD: float, L: float, G: float,
                              title: str) -> dict:
    """Old points rescored: needs `old_accepts` (verdict of the constraints of that time) next to the new ones."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    names = [c.name for c in constraints_for(config)]
    scored = scored.copy()
    scored["old_outcome"] = outcome_of(scored["old_accepts"], scored["built"])
    scored["verdict_changed"] = scored["old_accepts"].astype(bool) != scored["accepted"]
    scored["change"] = np.where(~scored["verdict_changed"], "",
                                np.where(scored["accepted"], "newly accepted", "newly rejected"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Constraints"
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=2, column=1, value=f"Configuration (N={config[0]}, T={config[1]});  OD={OD:g}  L={L:g}  G={G:g};  A in degrees")
    row = _write_block(ws, constraints_table(config), 4, "NEW constraints (rescored)")
    ws.cell(row=row, column=1, value="OLD constraints").font = Font(bold=True, size=12)
    ws.cell(row=row + 1, column=1, value=OLD_SPEC_TEXT)
    _style_sheet(ws, 5, freeze=False)

    ws = wb.create_sheet("Summary")
    new_sum = summary_table(scored, "outcome", "new constraints")
    old_sum = summary_table(scored, "old_outcome", "old constraints")
    both = pd.concat([old_sum, new_sum], ignore_index=True).sort_values(["subset", "spec"], kind="stable", ascending=[True, False])
    both = pd.concat([both[both["subset"] == "All points"], both[both["subset"] != "All points"]], ignore_index=True)
    row = _write_block(ws, both, 1, "Confusion matrix: old vs new constraints on the same points")
    changed = pd.DataFrame([{"verdict change": k, "points": int((scored["change"] == k).sum()),
                             "of which built": int(((scored["change"] == k) & scored["built"]).sum())}
                            for k in ("newly accepted", "newly rejected")])
    row = _write_block(ws, changed, row, "Points whose verdict changed")
    row = _write_block(ws, per_constraint_table(scored, config), row, "Each NEW constraint on its own")
    _write_block(ws, _legend(), row, "Legend")
    _style_sheet(ws, 2, freeze=False)

    front = ["row", "source", "label"]
    back = ["accepted", "built", "failed_features", "outcome", "old_accepts", "old_outcome", "change"]
    table = _point_columns(scored, names, front, back).rename(
        columns={"accepted": "new_constraints_accept", "built": "solidworks_built", "outcome": "new_outcome"})
    _data_sheet(wb, "All points", table, "no points")
    _data_sheet(wb, "Changed verdicts", table[scored["verdict_changed"].to_numpy()], "No verdict changed.")
    _data_sheet(wb, "False positives (new)", table[scored["outcome"] == "FP"], "No false positives.")
    _data_sheet(wb, "False negatives (new)", table[scored["outcome"] == "FN"], "No false negatives.")
    wb.save(path)
    return {"summary": both, "per_constraint": per_constraint_table(scored, config), "changed": changed}


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def _to_bool(col: pd.Series) -> pd.Series:
    if col.dtype == bool:
        return col
    return col.astype(str).str.strip().str.lower().map({"true": True, "false": False, "1": True, "0": False}).astype(bool)


def load_run(results_path: str, meta_path: str, allow_partial: bool = False) -> pd.DataFrame:
    """swgen_results.jsonl joined to the suite's meta CSV by row number (source / label)."""
    res = load_results(results_path)
    meta = pd.read_csv(meta_path)
    if len(res) < len(meta) and not allow_partial:
        raise ValueError(f"{results_path} has {len(res)} rows but {meta_path} lists {len(meta)}: the SwGen run is not "
                         f"finished (use --allow-partial to score what is there)")
    df = meta.merge(res, on="row", suffixes=("_meta", ""), how="inner")
    for c in PARAM_COLS:                                           # the rows must be the points that were requested
        bad = ~np.isclose(df[f"{c}_meta"].astype(float), df[c].astype(float))
        if bad.any():
            raise ValueError(f"result rows do not match the points file in column {c} (first row {int(df.loc[bad, 'row'].iloc[0])})")
    return df.drop(columns=[f"{c}_meta" for c in PARAM_COLS])


def load_old(csv_path: str) -> pd.DataFrame:
    df = pd.read_csv(csv_path)
    need = {*PARAM_COLS, "built", "current_spec_accepts"}
    if need - set(df.columns):
        raise ValueError(f"{csv_path} lacks column(s) {sorted(need - set(df.columns))}")
    df["built"] = _to_bool(df["built"])
    df["old_accepts"] = _to_bool(df["current_spec_accepts"])
    return df.rename(columns={"failed_features": "failed_features"})


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _print_tables(tables: dict) -> None:
    pd.set_option("display.width", 200)
    for name, t in tables.items():
        print(f"\n{name}\n{t.to_string(index=False)}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Excel reports: constraints vs SolidWorks rebuilds.")
    sub = parser.add_subparsers(dest="command", required=True)

    def common(p):
        p.add_argument("--config", default="6:1")
        p.add_argument("--OD", type=float, required=True)
        p.add_argument("--L", type=float, required=True)
        p.add_argument("--G", "--gap", type=float, default=0.0)
        p.add_argument("--out", required=True, help="output .xlsx")
        p.add_argument("--title", default=None)

    p_run = sub.add_parser("run", help="report on a fresh SwGen run")
    p_run.add_argument("--results", required=True, help="swgen_results.jsonl")
    p_run.add_argument("--meta", required=True, help="<points>.meta.csv written by validate_constraints.py suite")
    p_run.add_argument("--allow-partial", action="store_true")
    common(p_run)

    p_old = sub.add_parser("old", help="rescore an earlier run (CSV with built + current_spec_accepts) against the current constraints")
    p_old.add_argument("--old-csv", required=True)
    common(p_old)
    args = parser.parse_args(argv)

    try:
        (config,) = parse_configs(args.config)
        if args.command == "run":
            scored = score_points(load_run(args.results, args.meta, args.allow_partial), config, args.OD, args.L, args.G)
            title = args.title or f"Constraint validation vs SolidWorks rebuilds - {len(scored)} points"
            tables = write_run_workbook(args.out, scored, config, args.OD, args.L, args.G, title)
        else:
            scored = score_points(load_old(args.old_csv), config, args.OD, args.L, args.G)
            title = args.title or f"Earlier run ({len(scored)} points) rescored against the current constraints"
            tables = write_comparison_workbook(args.out, scored, config, args.OD, args.L, args.G, title)
    except (ValueError, NotImplementedError, FileNotFoundError) as exc:
        print(f"ERROR: {exc}")
        return 2
    _print_tables(tables)
    print(f"\nworkbook written to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
